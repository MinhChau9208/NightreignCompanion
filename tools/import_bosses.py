"""Regenerate data/bosses.json and data/boss_priors.json from the community
spreadsheets (kept out of git in data_source/).

    pip install openpyxl
    python tools/import_bosses.py "data_source/Nightreign Useful Info (Version 1.03.2) (+ DLC).xlsx" \
        "data_source/Nightlord Boss Pools.xlsx"

Stats come from the "Nightlord Stats {Solo,Duo,Trio}" and "Everdark Sovereign
Stats {Solo,Duo,Trio}" sheets; Duo/Trio only differ in health and poise. Pools
come from the two tables of the Boss Pools sheet (Night 1 first, Night 2 second).
Rerun after a game patch, then `go run ./cmd/nrc-cli validate`.
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
PARTY = ["Solo", "Duo", "Trio"]
DAMAGE = ["standard", "strike", "slash", "pierce", "magic", "fire", "lightning", "holy"]
STATUS = ["poison", "rot", "bleed", "frost", "sleep", "madness", "blight"]

# Columns of the stats sheets (0-based): name, game id, health, the eight
# damage negations (value column right of each merged header), status resists, poise.
COL_NAME, COL_ID, COL_HEALTH = 0, 2, 3
COL_DAMAGE = [6, 8, 10, 12, 14, 16, 18, 20]
COL_STATUS = [22, 23, 24, 25, 26, 27, 28]
COL_POISE = 29

NIGHTLORDS = [
    # id, name, expedition (None for DLC until named), pool column, row-name prefixes
    ("gladius", "Gladius, Beast of Night", "Tricephalos", "Tri", ["Gladius"]),
    ("adel", "Adel, Baron of Night", "Gaping Jaw", "GJ", ["Adel"]),
    ("gnoster", "Gnoster, Wisdom of Night", "Sentient Pest", "SP", ["Gnoster", "Faurtis", "Animus"]),
    ("maris", "Maris, Fathom of Night", "Augur", "Aug", ["Maris", "Augur"]),
    ("libra", "Libra, Creature of Night", "Equilibrious Beast", "EB", ["Libra"]),
    ("fulghor", "Fulghor, Champion of Nightglow", "Darkdrift Knight", "DDK", ["Fulghor"]),
    ("caligo", "Caligo, Miasma of Night", "Fissure in the Fog", "FitF", ["Caligo"]),
    ("heolstor", "Heolstor the Nightlord", "Night Aspect", "NA", ["The Shape of Night", "Heolstor"]),
    ("harmonia", "Weapon-Bequeathed Harmonia", None, None, ["Weapon-Bequeathed Harmonia"]),
    ("straghess", "Straghess", None, None, ["Traitorous Straghess", "Pure Impulse Straghess"]),
]
DLC = {"harmonia", "straghess"}

SUFFIX = re.compile(r"\s*\((phase \d+|single dog)\)\s*$", re.I)


def loc(s):
    # Boss names are proper nouns; the game keeps them in English in vi too.
    return {"en": s, "vi": s}


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", s.lower().replace("'", "")).strip("_")


def num(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def owner(name):
    for nl in NIGHTLORDS:
        if any(name.startswith(p) for p in nl[4]):
            return nl[0]
    raise SystemExit(f"stats row {name!r} matches no nightlord")


def read_rows(wb, sheet):
    rows = []
    for r in list(wb[sheet].iter_rows(values_only=True))[1:]:
        if not r[COL_NAME]:
            continue
        rows.append(r)
    return rows


def forms(wb, prefix):
    """Forms per nightlord for one difficulty, merging rows with equal stats
    (fight phases and Gladius' single dogs only repeat the same numbers)."""
    party = [read_rows(wb, f"{prefix} {p}") for p in PARTY]
    out = {}
    for solo, duo, trio in zip(*party):
        name = solo[COL_NAME].strip()
        if not (name == duo[COL_NAME].strip() == trio[COL_NAME].strip()):
            raise SystemExit(f"{prefix}: rows out of order at {name!r}")
        health = [num(r[COL_HEALTH]) for r in (solo, duo, trio)]
        form = {
            "name": SUFFIX.sub("", name),
            "gameIds": [int(solo[COL_ID])],
            "health": None if any(not isinstance(h, (int, float)) for h in health) else health,
            "poise": [round(r[COL_POISE], 1) for r in (solo, duo, trio)],
            "negation": {d: num(solo[c]) for d, c in zip(DAMAGE, COL_DAMAGE)},
            "resist": {},
            "immune": [],
        }
        for s, c in zip(STATUS, COL_STATUS):
            v = solo[c]
            if isinstance(v, str) and v.strip().lower() == "immune":
                form["immune"].append(s)
            else:
                form["resist"][s] = num(v)
        key = lambda f: json.dumps({k: f[k] for k in ("health", "poise", "negation", "resist", "immune")})
        lst = out.setdefault(owner(name), [])
        same = next((f for f in lst if key(f) == key(form)), None)
        if same is None:
            lst.append(form)
            continue
        if int(solo[COL_ID]) not in same["gameIds"]:
            same["gameIds"].append(int(solo[COL_ID]))
        if form["name"] != same["name"]:
            same["name"] += " / " + form["name"]
    return out


def pools(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    tables, cols = [], None
    for r in ws.iter_rows(values_only=True):
        if r[0] == "Bosses":
            cols = list(r[1:])
            tables.append([])
        elif r[0] and cols:
            tables[-1].append((r[0].strip(), {c: bool(v) for c, v in zip(cols, r[1:])}))
    if len(tables) != 2:
        raise SystemExit(f"expected 2 pool tables (Night 1, Night 2), got {len(tables)}")
    return tables


def main(stats_path, pools_path):
    wb = openpyxl.load_workbook(stats_path, data_only=True)
    normal = forms(wb, "Nightlord Stats")
    everdark = forms(wb, "Everdark Sovereign Stats")
    night1, night2 = pools(pools_path)

    bosses = []
    for nid, name, exp, _, _ in NIGHTLORDS:
        b = {"id": nid, "kind": "nightlord", "name": loc(name)}
        if exp:
            b["expedition"] = loc(exp)
        b["dlc"] = nid in DLC
        b["stats"] = {"normal": normal.get(nid, []), "everdark": everdark.get(nid, [])}
        b["verified"] = True
        bosses.append(b)
    for night, table in ((1, night1), (2, night2)):
        for name, _ in table:
            bosses.append({"id": slug(name), "kind": "night_boss", "name": loc(name), "night": night, "verified": True})

    priors = {
        "verified": True,
        "note": "Night-boss pools per Nightlord (Deep of Night, hidden Nightlord), from 'Nightlord Boss Pools.xlsx'. "
        "DLC Nightlords have no pool yet.",
        "pools": [],
    }
    for nid, _, _, col, _ in NIGHTLORDS:
        if col is None:
            continue
        priors["pools"].append({
            "nightlord": nid,
            "night1": [slug(n) for n, has in night1 if has[col]],
            "night2": [slug(n) for n, has in night2 if has[col]],
        })

    write(ROOT / "data" / "bosses.json", bosses)
    write(ROOT / "data" / "boss_priors.json", priors)


def write(path, v):
    path.write_text(json.dumps(v, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print("wrote", path.relative_to(ROOT))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
